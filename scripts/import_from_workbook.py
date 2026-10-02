"""
Run the three D12 Ball import scripts on the sheet's workbook, for when
the sheet's CSV export cannot be fetched.

The import scripts read each tab as a CSV, by default straight from the
sheet's `export?format=csv&gid=...` URLs. A cloud session whose network
policy refuses `docs.google.com`, or the `*-sheets.googleusercontent.com`
host that export redirects to, cannot reach them -- but the Google Drive
connector can export the whole workbook as `.xlsx`. This takes that
workbook, writes the five tabs the importers read as CSVs, and runs
`import_d12ball_players.py`, `import_d12ball_species.py` and
`import_d12ball_maneuvers.py` on them. The importers are the same either
way, and the data they write names the sheet's own URLs as its source,
so a run from here writes what a run off the sheet would.

The workbook may be the `.xlsx` itself or the JSON the Drive connector's
`download_file_content` saves when its result is too large to show
(`{"content": <base64>, ...}`), exported as
`application/vnd.openxmlformats-officedocument.spreadsheetml.sheet`.

    python3 scripts/import_from_workbook.py <workbook> [--only players]
        [--data-version maneuvers=15]

Each file keeps the `data_version` it has unless `--data-version` gives
it a new one -- the importers' own default is 1, which would reset it.

Needs `openpyxl` (`pip install openpyxl`), which the bot does not.
See docs/design/rules-and-data.md.
"""

from __future__ import annotations

import argparse
import base64
import csv
import json
import subprocess
import sys
import tempfile
from pathlib import Path

SCRIPTS = Path(__file__).resolve().parent
DATA = SCRIPTS.parent / "d12ball" / "data"

#: The tab each importer argument reads, by title. Titles rather than
#: gids because a workbook file carries titles; the importers' own
#: `DEFAULT_*` URLs name the same tabs by gid. Matched with surrounding
#: whitespace ignored -- `player cards ` and `advanced_abilities ` end in
#: a space on the sheet.
IMPORTS = {
    "players": (
        "players.json",
        "import_d12ball_players.py",
        {
            "--source": "player cards",
            "--abilities": "basic_abilities",
            "--advanced": "advanced_abilities",
        },
    ),
    "species": (
        "species.json",
        "import_d12ball_species.py",
        {"--source": "spec_abilities"},
    ),
    "maneuvers": (
        "maneuvers.json",
        "import_d12ball_maneuvers.py",
        {"--source": "maneuvers"},
    ),
}


def workbook_bytes(path: Path) -> bytes:
    """The `.xlsx` itself, or the one inside a connector's JSON dump."""
    raw = path.read_bytes()
    if raw[:2] == b"PK":
        return raw
    return base64.b64decode(json.loads(raw)["content"])


def write_tab(workbook, title: str, out: Path) -> None:
    """One tab as the CSV the sheet's export would have served."""
    sheets = {name.strip(): name for name in workbook.sheetnames}
    if title not in sheets:
        raise SystemExit(
            f"The workbook has no tab {title!r}: it has "
            + ", ".join(repr(name) for name in workbook.sheetnames)
        )
    with out.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        for row in workbook[sheets[title]].iter_rows(values_only=True):
            writer.writerow([cell_text(value) for value in row])


def current_version(output: str) -> int:
    """
    The `data_version` the file already carries, so a run here never
    resets it to the importers' default of 1. Bumping it for a change
    is the caller's, as it is on a run off the sheet (`--data-version`).
    """
    try:
        return int(json.loads((DATA / output).read_text())["data_version"])
    except (OSError, ValueError, KeyError, TypeError):
        return 1


def cell_text(value) -> str:
    # The export writes a whole number as "6", not "6.0".
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def _version_pair(text: str) -> tuple[str, int]:
    name, _, number = text.partition("=")
    if name not in IMPORTS or not number.isdigit():
        raise argparse.ArgumentTypeError(
            f"expected NAME=N with NAME one of {', '.join(IMPORTS)}"
        )
    return name, int(number)


def main() -> None:
    parser = argparse.ArgumentParser(
        description=(
            "Run the D12 Ball import scripts on the sheet's workbook "
            "rather than its CSV export."
        ),
    )
    parser.add_argument(
        "workbook",
        type=Path,
        help="The .xlsx, or the Drive connector's saved JSON around it.",
    )
    parser.add_argument(
        "--only",
        choices=sorted(IMPORTS),
        action="append",
        help="Run just these importers (default: all three).",
    )
    parser.add_argument(
        "--data-version",
        action="append",
        type=_version_pair,
        metavar="NAME=N",
        help=(
            "A new data_version for one importer's file, e.g. "
            "maneuvers=15; every other file keeps the one it has."
        ),
    )
    args = parser.parse_args()
    args.data_version = dict(args.data_version or ())

    try:
        from openpyxl import load_workbook
    except ImportError:
        raise SystemExit("This needs openpyxl: pip install openpyxl")

    with tempfile.TemporaryDirectory() as scratch:
        scratch = Path(scratch)
        xlsx = scratch / "sheet.xlsx"
        xlsx.write_bytes(workbook_bytes(args.workbook))
        workbook = load_workbook(xlsx, data_only=True)
        for name in args.only or IMPORTS:
            output, script, tabs = IMPORTS[name]
            version = (
                args.data_version.get(name)
                if args.data_version and name in args.data_version
                else current_version(output)
            )
            command = [
                sys.executable, str(SCRIPTS / script),
                "--data-version", str(version),
            ]
            for flag, title in tabs.items():
                out = scratch / f"{title}.csv"
                write_tab(workbook, title, out)
                command += [flag, str(out)]
            subprocess.run(command, check=True)


if __name__ == "__main__":
    main()
